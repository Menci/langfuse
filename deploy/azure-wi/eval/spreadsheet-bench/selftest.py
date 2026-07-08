"""Self-test for the vendored SpreadsheetBench objective gate.

Proves the harness is wired correctly *without* running any model, by driving
the vendored ``compare_workbooks`` (from ``evaluation/evaluation.py``) three
ways:

  1. Synthetic micro-test (deterministic): two in-memory workbooks that are
     identical compare equal; changing one cell in the checked range compares
     unequal. This directly validates cell comparison + range parsing.
  2. Real-data sanity: for a sample of Verified-400 tasks whose answer range is
     well-formed, golden-vs-golden must compare **equal** (True).
  3. Real-data discriminative: among those, golden-vs-init must compare
     **unequal** (False) for most tasks — proving the harness detects genuine
     differences between the unmodified input and the golden answer.

Some Verified-400 tasks carry non-standard ``answer_position`` values (whole
columns like ``A:G``, or malformed tokens); like the upstream evaluator, we
treat comparison exceptions defensively rather than crashing.

Usage:
    python eval/spreadsheet-bench/selftest.py [--sample 40] [--all]

Pure openpyxl (no LibreOffice needed): comparisons read cached values with
``data_only=True`` and golden/init files already carry cached values.
"""

from __future__ import annotations

import argparse
import contextlib
import importlib.util
import io
import sys
import tempfile
from pathlib import Path

import openpyxl

_HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(_HERE))
import download_data as dd  # noqa: E402  (shared dataset helpers)


def _load_vendored_evaluation():
    """Import the vendored evaluation.py by path (its __main__ guard is safe)."""
    eval_py = _HERE / "evaluation" / "evaluation.py"
    spec = importlib.util.spec_from_file_location("ssb_evaluation", eval_py)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def safe_compare(compare_workbooks, gt, proc, itype, ranges):
    """Return True/False, or None if the comparison raised (e.g. malformed
    range). Suppresses the vendored function's per-cell stdout chatter."""
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            ok, _ = compare_workbooks(str(gt), str(proc), itype, ranges)
        return bool(ok)
    except Exception:
        return None


def _synthetic_check(compare_workbooks) -> list[str]:
    """Return a list of failure messages (empty == pass)."""
    failures: list[str] = []
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        gt, same, diff = tmp / "gt.xlsx", tmp / "same.xlsx", tmp / "diff.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        for row, vals in enumerate([["a", 1, 2.5], ["b", 3, 4.0]], start=1):
            for col, v in enumerate(vals, start=1):
                ws.cell(row=row, column=col, value=v)
        wb.save(gt)
        wb.save(same)
        ws.cell(row=1, column=2, value=999)  # perturb B1 (inside A1:C2)
        wb.save(diff)

        if safe_compare(compare_workbooks, gt, same, "Cell-Level Manipulation", "Sheet1!A1:C2") is not True:
            failures.append("identical workbooks compared UNEQUAL (expected equal)")
        if safe_compare(compare_workbooks, gt, diff, "Cell-Level Manipulation", "Sheet1!A1:C2") is not False:
            failures.append("one-cell-different workbooks compared EQUAL (expected unequal)")
    return failures


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sample", type=int, default=40, help="Number of tasks to sample")
    parser.add_argument("--all", action="store_true", help="Use all tasks")
    args = parser.parse_args()

    ev = _load_vendored_evaluation()
    compare = ev.compare_workbooks

    # ── Check 1: synthetic ────────────────────────────────────────────────
    print("== Check 1: synthetic micro-test ==")
    syn_failures = _synthetic_check(compare)
    for f in syn_failures:
        print(f"  FAIL: {f}")
    if not syn_failures:
        print("  OK: identical->equal, one-cell-diff->unequal")

    # ── Real data ─────────────────────────────────────────────────────────
    root = dd.ensure_dataset()
    records = dd.load_records(root)
    complete = [r for r in records if list(dd.iter_test_cases(r, root))]
    complete.sort(key=lambda r: str(r["id"]))
    sample = complete if args.all else complete[: args.sample]

    gg_equal, gg_mismatch, gg_error = 0, 0, 0
    gi_differ, gi_same = 0, 0
    gg_error_ids, gg_mismatch_ids, gi_same_ids = [], [], []

    for rec in sample:
        idx, init_path, golden_path = next(dd.iter_test_cases(rec, root))
        ranges = dd.answer_ranges(rec)
        itype = rec.get("instruction_type", "")

        gg = safe_compare(compare, golden_path, golden_path, itype, ranges)
        if gg is None:
            gg_error += 1
            gg_error_ids.append(str(rec["id"]))
            continue  # unresolvable range; skip discriminative check
        if gg is not True:
            gg_mismatch += 1
            gg_mismatch_ids.append(str(rec["id"]))
            continue
        gg_equal += 1

        gi = safe_compare(compare, golden_path, init_path, itype, ranges)
        if gi is False:
            gi_differ += 1
        else:  # True (init already matches) or None (treat as non-discriminative)
            gi_same += 1
            gi_same_ids.append(str(rec["id"]))

    n = len(sample)
    print(f"\nDataset: {len(records)} records, {len(complete)} with test cases; checking {n}.")
    print("\n== Check 2: golden-vs-golden (expect equal) ==")
    print(f"  equal: {gg_equal}/{n} | mismatch: {gg_mismatch} | unresolvable-range: {gg_error}")
    if gg_mismatch_ids:
        print(f"    mismatch ids: {gg_mismatch_ids[:10]}")
    if gg_error_ids:
        print(f"    unresolvable-range ids (non-standard answer_position): {gg_error_ids[:10]}")

    resolvable = gg_equal
    differ_rate = gi_differ / resolvable if resolvable else 0.0
    print("\n== Check 3: golden-vs-init among resolvable (expect differ) ==")
    print(f"  differ: {gi_differ}/{resolvable} ({differ_rate:.0%}) | init-already-matches: {gi_same}")
    if gi_same_ids:
        print(f"    init-already-matches ids: {gi_same_ids[:10]}")

    # ── Verdict ───────────────────────────────────────────────────────────
    ok = (
        not syn_failures
        and resolvable >= 0.5 * n      # enough tasks have well-formed ranges
        and gg_mismatch == 0           # golden must always match itself
        and differ_rate >= 0.80        # harness detects real differences
    )
    print("\n" + "=" * 48)
    print("SELF-TEST PASSED" if ok else "SELF-TEST FAILED")
    print("=" * 48)
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
